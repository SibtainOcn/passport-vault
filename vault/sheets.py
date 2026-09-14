import csv, io, zipfile
from datetime import date, datetime
from openpyxl import load_workbook, Workbook as ExcelWorkbook
from openpyxl.styles import PatternFill, Font
from .fields import FIELDS, normalize
MAX_ROWS=1000000
MAX_COLS=40

def scalar(v):
    if isinstance(v,(datetime,date)): return v.strftime('%Y-%m-%d')
    if v is None: return ''
    if isinstance(v,bool): return str(v)
    s=str(v)
    if len(s)>2000: raise ValueError('Cell exceeds 2000 characters.')
    return s

def parse_book(data, filename):
    if filename.lower().endswith('.csv'):
        try: text=data.decode('utf-8-sig')
        except UnicodeDecodeError: raise ValueError('CSV must be UTF-8. Save as CSV UTF-8 from Excel.')
        reader=csv.reader(io.StringIO(text)); rows=[]
        for i,row in enumerate(reader):
            if i>=MAX_ROWS+20 or len(row)>MAX_COLS: raise ValueError('Workbook limit: 1000000 data rows and 40 columns.')
            rows.append([scalar(v) for v in row])
        sheets={'CSV':rows}
    elif filename.lower().endswith('.xlsx'):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if len(z.infolist())>3000 or sum(f.file_size for f in z.infolist())>256*1024*1024: raise ValueError('Workbook expands beyond safe limits.')
                if any('vbaproject' in n.lower() or n.startswith('xl/externalLinks/') for n in z.namelist()): raise ValueError('Macros and external links are not supported.')
            wb=load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
            if len(wb.worksheets)>30: raise ValueError('At most 30 sheets are allowed.')
            sheets={}; total=0
            try:
                for ws in wb.worksheets:
                    if ws.max_column and ws.max_column>MAX_COLS: raise ValueError('Maximum 40 columns per sheet.')
                    if ws.max_row and ws.max_row>MAX_ROWS+20: raise ValueError('Maximum 1000000 data rows plus headers.')
                    rows=[]
                    for i,row in enumerate(ws.iter_rows()):
                        total+=1
                        if total>MAX_ROWS+100 or i>MAX_ROWS+20 or len(row)>MAX_COLS: raise ValueError('Workbook row/column limit exceeded.')
                        vals=[]
                        for cell in row:
                            if cell.data_type=='f': raise ValueError('Formula cells are not accepted. Upload a values-only copy.')
                            if cell.data_type=='e': raise ValueError('Excel contains an error cell. Correct it before upload.')
                            vals.append(scalar(cell.value))
                        rows.append(vals)
                    sheets[ws.title]=rows
            finally: wb.close()
        except ValueError: raise
        except Exception: raise ValueError('Corrupted, encrypted or unsupported Excel workbook.')
    else: raise ValueError('Use .xlsx or UTF-8 .csv; .xls and .xlsm are not supported.')
    if not sheets or not any(sheets.values()): raise ValueError('Workbook is empty.')
    return sheets

def safe_cell(value):
    text=str(value if value is not None else '')
    return "'"+text if text.lstrip().startswith(('=','+','-','@','\t','\r','\n')) else text

def export_report(report,kind):
    headers=['Sheet','Row','Passport number','Result','Reason','Duplicate input','Record ID','Record revision']
    selected=report['fields']
    for f in selected: headers += [f+' (Excel)',f+' (Passport)',f+' result']
    rows=[]
    for row in report['rows']:
        vals=[row['sheet'],row['row'],row['passport_number'],row['status'],row['reason'],row['duplicate_input'],row.get('document_id',''),row.get('revision','')]
        for f in selected:
            item=row['comparisons'].get(f,{})
            vals += [item.get('excel',''),item.get('passport',''),item.get('status','Not compared')]
        rows.append([safe_cell(x) for x in vals])
    if kind=='csv':
        out=io.StringIO(newline=''); w=csv.writer(out); w.writerow(headers); w.writerows(rows)
        return ('\ufeff'+out.getvalue()).encode('utf-8'),'text/csv; charset=utf-8'
    wb=ExcelWorkbook(); ws=wb.active; ws.title='Comparison'; ws.append(headers)
    colors={'Match':'E2F4EA','Mismatch':'FCE4E4','Needs Review':'FFF1CC','Not Found':'E9ECF2','Invalid Row':'FCE4E4'}
    for row,vals in zip(report['rows'],rows):
        ws.append(vals)
        for c in ws[ws.max_row]:
            c.data_type='s'; c.fill=PatternFill('solid',fgColor=colors.get(row['status'],'FFFFFF'))
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='172E46')
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    for column in ws.columns: ws.column_dimensions[column[0].column_letter].width=24
    meta=wb.create_sheet('Report details')
    for k in ('created','rules_version','date_format','scope','warning'): meta.append([k,safe_cell(report.get(k,''))])
    meta.column_dimensions['A'].width=24; meta.column_dimensions['B'].width=90
    out=io.BytesIO(); wb.save(out); wb.close()
    return out.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
