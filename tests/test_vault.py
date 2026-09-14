import io, json, os, sys, time, zipfile
from unittest.mock import patch
from datetime import timedelta
from PIL import Image, ImageFont

def _mono_font(size):
    """Return a monospace TrueType font, trying common OS locations."""
    candidates = [
        '/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf',           # Linux (Debian/Ubuntu)
        '/usr/share/fonts/dejavu-sans-mono-fonts/DejaVuSansMono.ttf',    # Linux (Fedora/RHEL)
        '/System/Library/Fonts/Menlo.ttc',                               # macOS
    ]
    # Windows: check common font directory
    windir = os.environ.get('WINDIR', r'C:\Windows')
    candidates.append(os.path.join(windir, 'Fonts', 'consola.ttf'))      # Consolas
    candidates.append(os.path.join(windir, 'Fonts', 'cour.ttf'))         # Courier New
    for path in candidates:
        if os.path.isfile(path):
            return ImageFont.truetype(path, size)
    # Last resort: Pillow default bitmap font (no sizing support, but won't crash)
    return ImageFont.load_default()
import pyotp
from cryptography.exceptions import InvalidTag
from django.test import TestCase, SimpleTestCase, Client
from django.contrib.auth import get_user_model
from django.contrib.auth.hashers import make_password
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.utils import timezone
from openpyxl import Workbook as XL, load_workbook
from vault.crypto import seal, unseal, pack, unpack, blind
from vault.fields import normalize, validate_fields
from vault.mrz import parse_td3, digit
from vault.models import Document, OwnerSecurity, Page, Source, Workbook, Report, Audit, LoginGuard
from vault.comparison import compare
from vault.sheets import parse_book, export_report
from vault.ocr import extract

A='P<UTOERIKSSON<<ANNA<MARIA<<<<<<<<<<<<<<<<<<<'
B='L898902C36UTO7408122F1204159ZE184226B<<<<<10'
# ICAO sample MRZ first line is padded explicitly to 44 characters.
A=A.ljust(44,'<')
class PureTests(SimpleTestCase):
 def test_encryption_random_nonce_and_aad(self):
  a,b=seal(b'sensitive','one'),seal(b'sensitive','one')
  self.assertNotEqual(a,b); self.assertEqual(unseal(a,'one'),b'sensitive')
  with self.assertRaises(InvalidTag):unseal(a,'two')
  with self.assertRaises(InvalidTag):unseal(a[:-1]+bytes([a[-1]^1]),'one')
 def test_mrz_checks(self):
  p=parse_td3(A,B);self.assertTrue(p['valid']);self.assertEqual(p['fields']['passport_number'],'L898902C3')
  self.assertFalse(parse_td3(A,'M'+B[1:])['valid'])
 def test_date_rules(self):
  self.assertEqual(normalize('dob','03/04/2000','DMY'),'2000-04-03')
  self.assertEqual(normalize('dob','03/04/2000','MDY'),'2000-03-04')
  with self.assertRaises(ValueError): normalize('dob','31/02/2000')
 def test_no_silent_passport_replacement(self):
  self.assertNotEqual(normalize('passport_number','A0123456'),normalize('passport_number','AO123456'))
  with self.assertRaises(ValueError): normalize('passport_number','A 123456')
 def test_formula_rejected(self):
  w=XL();w.active.append(['Passport','Name']);w.active.append(['A1234567','=1+1']);b=io.BytesIO();w.save(b)
  with self.assertRaises(ValueError): parse_book(b.getvalue(),'test.xlsx')
 def test_csv_and_utf8(self):
  self.assertEqual(parse_book(b'Passport,Name\nA1234567,TEST USER','test.csv')['CSV'][1][0],'A1234567')
  with self.assertRaises(ValueError):parse_book(b'\xff','x.csv')
 def test_wrong_excel_extension(self):
  with self.assertRaises(ValueError):parse_book(b'data','x.xlsm')
 def test_zip_expansion_limit(self):
  b=io.BytesIO()
  with zipfile.ZipFile(b,'w',zipfile.ZIP_DEFLATED) as z:z.writestr('large',b'0'*(41*1024*1024))
  with self.assertRaises(ValueError):parse_book(b.getvalue(),'x.xlsx')
 def test_field_dates_order(self):
  with self.assertRaises(ValueError):validate_fields({'passport_number':'A1234567','full_name':'TEST USER','date_of_issue':'2030-01-01','date_of_expiry':'2020-01-01'})

class VaultTests(TestCase):
 def setUp(self):
  self.user=get_user_model().objects.create_user(username='owner',password='test-password-long-123')
  self.client.force_login(self.user);s=self.client.session;s['mfa']=True;s.save()
 def doc(self,number='A1234567',status='approved',**fields):
  d=Document(owner=self.user,status=status,content_index=blind(str(time.time_ns())))
  data={'name':'private-file.png','fields':{'passport_number':number,'full_name':'TEST USER','dob':'2000-04-03','issuing_country':'IND',**fields},'warnings':[]}
  d.payload=pack(data,'document:'+str(d.id));d.passport_index=blind(number);d.save();return d
 def comparison(self,number='A1234567',name='TEST USER',dob='03/04/2000'):
  return compare(self.user,{'Sheet1':[['Passport','Name','DOB'],[number,name,dob]]},{'Sheet1':1},{'passport_number':0,'full_name':1,'dob':2},'DMY')
 def post(self,url,data):return self.client.post(url,json.dumps(data),content_type='application/json')
 def test_comparison_match_and_mismatch(self):
  self.doc();self.assertEqual(self.comparison()['rows'][0]['status'],'Match')
  row=self.comparison(name='DIFFERENT')['rows'][0];self.assertEqual(row['status'],'Mismatch');self.assertEqual(row['comparisons']['full_name']['status'],'Mismatch')
 def test_notfound_and_pending(self):
  self.assertEqual(self.comparison()['rows'][0]['status'],'Not Found')
  self.doc('B1234567','queued');self.assertEqual(self.comparison()['rows'][0]['status'],'Needs Review')
 def test_duplicate_candidates_and_country(self):
  self.doc();self.doc();self.assertEqual(self.comparison()['rows'][0]['status'],'Needs Review')
 def test_missing_dob_not_match(self):
  self.doc(dob='');self.assertEqual(self.comparison()['rows'][0]['status'],'Needs Review')
 def test_unapproved_not_match(self):
  self.doc(status='review');self.assertEqual(self.comparison()['rows'][0]['status'],'Needs Review')
 def test_invalid_input_row(self):
  self.assertEqual(self.comparison(number='')['rows'][0]['status'],'Invalid Row')
 def test_other_owner_not_matched(self):
  d=self.doc();other=get_user_model().objects.create_user(username='other');d.owner=other;d.save()
  self.assertEqual(self.comparison()['rows'][0]['status'],'Not Found')
 def test_no_access_without_mfa(self):
  s=self.client.session;s.pop('mfa');s.save();self.assertEqual(self.client.get('/api/documents').status_code,401)
 def test_anonymous_blocked(self):
  self.client.logout();self.assertEqual(self.client.get('/api/documents').status_code,401)
 def test_object_authorization(self):
  d=self.doc();other=get_user_model().objects.create_user(username='other');d.owner=other;d.save()
  self.assertEqual(self.client.get('/api/documents/'+str(d.id)).status_code,404)
  self.assertEqual(self.post('/api/documents/'+str(d.id)+'/delete',{'confirm':'DELETE'}).status_code,404)
 def test_csrf_enforced(self):
  c=Client(enforce_csrf_checks=True);c.force_login(self.user);s=c.session;s['mfa']=True;s.save()
  self.assertEqual(c.post('/api/logout',data='{}',content_type='application/json').status_code,403)
 def test_headers(self):
  r=self.client.get('/api/session');self.assertEqual(r['Cache-Control'],'no-store');self.assertIn("frame-ancestors 'none'",r['Content-Security-Policy'])
 def test_upload_encryption_and_duplicate(self):
  b=io.BytesIO();Image.new('RGB',(20,20),'white').save(b,format='PNG');raw=b.getvalue()
  response=self.client.post('/api/upload',{'files':SimpleUploadedFile('name.png',raw),'language':'eng'})
  self.assertEqual(response.status_code,201);s=Source.objects.get();self.assertNotEqual(bytes(s.data),raw);self.assertEqual(unseal(s.data,'source:'+str(s.id)),raw)
  self.assertEqual(self.client.post('/api/upload',{'files':SimpleUploadedFile('name.png',raw)}).status_code,409)
 def test_wrong_file_rejected(self):
  self.assertEqual(self.client.post('/api/upload',{'files':SimpleUploadedFile('x.pdf',b'not a pdf')}).status_code,400)
 def test_revision_conflict_and_history(self):
  d=self.doc(status='review');data=unpack(d.payload,'document:'+str(d.id));url='/api/documents/'+str(d.id)+'/review'
  self.assertEqual(self.post(url,{'fields':data['fields'],'revision':2,'verified':True}).status_code,409)
  self.assertEqual(self.post(url,{'fields':data['fields'],'revision':0,'verified':True}).status_code,200)
  d.refresh_from_db();self.assertEqual(d.revision,1);self.assertEqual(d.revisions.count(),1)
 def test_multiple_passports_cannot_approve(self):
  d=self.doc(status='review');data=unpack(d.payload,'document:'+str(d.id));data['multiple_identities']=True;d.payload=pack(data,'document:'+str(d.id));d.save()
  self.assertEqual(self.post('/api/documents/'+str(d.id)+'/review',{'fields':data['fields'],'revision':0,'verified':True}).status_code,400)
 def test_report_formula_safety(self):
  self.doc();report=self.comparison(name='=HYPERLINK("https://example.com")')
  csvdata,_=export_report(report,'csv');self.assertIn(b"'=HYPERLINK",csvdata)
  xlsx,_=export_report(report,'xlsx');wb=load_workbook(io.BytesIO(xlsx));self.assertFalse(any(c.data_type=='f' for row in wb.active for c in row));wb.close()
 def test_report_snapshot_persists(self):
  d=self.doc();report=self.comparison();r=Report(owner=self.user);r.data=pack(report,'report:'+str(r.id));r.save();d.delete()
  self.assertEqual(unpack(r.data,'report:'+str(r.id))['rows'][0]['status'],'Match')
 def test_audit_chain(self):
  self.client.get('/api/documents');call_command('maintenance',stdout=io.StringIO())
  Audit.objects.update(action='modified')
  from django.core.management.base import CommandError
  with self.assertRaises(CommandError):call_command('maintenance',stdout=io.StringIO())
 def test_exact_search(self):
  self.doc('A1234567');self.doc('B1234567');r=self.client.get('/api/documents?passport=A1234567').json();self.assertEqual(r['total'],1)
 def test_mfa_replay_and_recovery(self):
  self.client.logout();secret=pyotp.random_base32();code=pyotp.TOTP(secret).now()
  OwnerSecurity.objects.create(user=self.user,secret=seal(secret.encode(),'mfa:'+str(self.user.id)),recovery=[make_password('recovery-once')])
  payload={'username':'owner','password':'test-password-long-123','code':code}
  self.assertEqual(self.post('/api/login',payload).status_code,200);self.client.logout()
  self.assertEqual(self.post('/api/login',payload).status_code,401)
  payload['code']='recovery-once';self.assertEqual(self.post('/api/login',payload).status_code,200);self.client.logout()
  self.assertEqual(self.post('/api/login',payload).status_code,401)
 def test_rate_limit(self):
  self.client.logout()
  for _ in range(5):self.post('/api/login',{'username':'bad','password':'bad','code':'123456'})
  self.assertEqual(self.post('/api/login',{}).status_code,429)
 def test_interrupted_worker(self):
  d=self.doc(status='processing');d.lease_until=timezone.now()-timedelta(seconds=1);d.save()
  call_command('worker',once=True);d.refresh_from_db();self.assertEqual(d.status,'failed')

class OCRSmoke(SimpleTestCase):
 def test_actual_tesseract_on_synthetic_document(self):
  from PIL import ImageDraw
  image=Image.new('RGB',(1800,1200),'white');draw=ImageDraw.Draw(image)
  font=_mono_font(38)
  lines=['SYNTHETIC TEST - NOT A TRAVEL DOCUMENT','Passport Number: A1234567','Full Name: TEST USER','Date of Birth: 03/04/2000','Place of Birth: TEST CITY']
  for i,line in enumerate(lines):draw.text((50,60+i*100),line,font=font,fill='black')
  buffer=io.BytesIO();image.save(buffer,format='PNG')
  result=extract([(buffer.getvalue(),'png')])
  self.assertEqual(result['fields']['passport_number'],'A1234567');self.assertEqual(len(result['previews']),1)

class WorkflowTests(TestCase):
 setUp = VaultTests.setUp
 post = VaultTests.post
 # End-to-end API workflow with real local OCR and isolated child processes.
 def test_upload_worker_review_compare_export(self):
  from PIL import ImageDraw
  image=Image.new('RGB',(1600,1000),'white');draw=ImageDraw.Draw(image);font=_mono_font(36)
  for i,line in enumerate(['SYNTHETIC TEST - NOT A PASSPORT','Passport Number: T1234567','Full Name: TEST PERSON','Date of Birth: 01/02/2000']):draw.text((50,50+i*110),line,font=font,fill='black')
  raw=io.BytesIO();image.save(raw,format='PNG')
  upload=self.client.post('/api/upload',{'files':SimpleUploadedFile('synthetic.png',raw.getvalue())});self.assertEqual(upload.status_code,201)
  pk=upload.json()['id'];call_command('worker',once=True)
  detail=self.client.get('/api/documents/'+pk).json();self.assertEqual(detail['status'],'review',detail);self.assertEqual(len(detail['pages']),1)
  preview=self.client.get('/api/pages/'+detail['pages'][0]['id']);self.assertEqual(preview.status_code,200);self.assertEqual(preview['Content-Type'],'image/jpeg')
  values=detail['fields'];values.update(passport_number='T1234567',full_name='TEST PERSON',dob='2000-02-01')
  self.assertEqual(self.post('/api/documents/'+pk+'/review',{'verified':True,'revision':0,'fields':values}).status_code,200)
  workbook=self.client.post('/api/workbooks',{'file':SimpleUploadedFile('synthetic.csv',b'Passport,Name,DOB\nT1234567,TEST PERSON,01/02/2000\n')});self.assertEqual(workbook.status_code,200,workbook.content)
  report=self.post('/api/compare',{'workbook':workbook.json()['id'],'sheets':{'CSV':1},'mapping':{'passport_number':0,'full_name':1,'dob':2},'date_format':'DMY'});self.assertEqual(report.status_code,200,report.content);self.assertEqual(report.json()['rows'][0]['status'],'Match')
  self.assertEqual(self.client.get('/api/reports/'+report.json()['id']+'/export/xlsx').status_code,200)
 def test_corrupt_pdf_worker_fails(self):
  response=self.client.post('/api/upload',{'files':SimpleUploadedFile('bad.pdf',b'%PDF- broken')});self.assertEqual(response.status_code,201)
  call_command('worker',once=True);doc=Document.objects.get(pk=response.json()['id']);self.assertEqual(doc.status,'failed');self.assertEqual(doc.error_code,'PDF_CORRUPT_OR_PASSWORD')
